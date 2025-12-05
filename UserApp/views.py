from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import logout, authenticate
from django.contrib.auth.views import LoginView
from django.contrib.auth.decorators import login_required
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy
from django.utils.decorators import method_decorator
from django.db.models import Q
from django.contrib import messages
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
from .forms import UserRegisterForm, ClientForm, LivreurForm
from .models import User

# Create your views here.
def register(req):
	if req.method == "POST":
		form = UserRegisterForm(req.POST)
		if form.is_valid():
			form.save()
			return redirect("login")
	else:
		form = UserRegisterForm()
	return render(req, "user/register.html", {'form': form})


class RoleLoginView(LoginView):
	template_name = 'user/login.html'

	def form_invalid(self, form):
		# Check if user exists but is inactive
		username = form.cleaned_data.get('username')
		password = form.cleaned_data.get('password')
		
		if username and password:
			try:
				user = User.objects.get(username=username)
				# Check password is correct but account is inactive
				if user.check_password(password) and not user.is_active:
					form.add_error(None, 'Your account has been deactivated. Please contact support.')
					return super().form_invalid(form)
			except User.DoesNotExist:
				pass
		
		return super().form_invalid(form)

	def get_success_url(self):
		user = self.request.user
		if user.role == 'gerant':
			return '/gerant/dashboard/'
		elif user.role == 'livreur':
			return '/livreur/home/'
		else:
			return '/client/home/'


def logout_view(req):
	logout(req)
	return redirect("home")


def client_home(request):
    """Client home page with restaurant menu integrated from islem"""
    # Fetch dishes by category from islem's API
    dishes_by_category = []
    try:
        import requests
        response = requests.get('http://127.0.0.1:8000/api/dishes/by-category/', timeout=2)
        if response.status_code == 200:
            dishes_by_category = response.json().get('categories', [])
    except:
        # If islem server is not running or API doesn't exist, show empty
        dishes_by_category = []
    
    context = {
        'dishes_by_category': dishes_by_category,
    }
    return render(request, 'client_home.html', context)

def livreur_home(request):
    return render(request, 'livreur_home.html')

def home(request):
	"""Home page with restaurant menu integrated from islem"""
	# Fetch dishes by category from islem's API
	dishes_by_category = []
	try:
		import requests
		response = requests.get('http://127.0.0.1:8000/api/dishes/by-category/', timeout=2)
		if response.status_code == 200:
			dishes_by_category = response.json().get('categories', [])
	except:
		# If islem server is not running or API doesn't exist, show empty
		dishes_by_category = []
	
	context = {
		'dishes_by_category': dishes_by_category,
		'is_authenticated': request.user.is_authenticated,
		'user_role': request.user.role if request.user.is_authenticated else None,
	}
	return render(request, 'home.html', context)


# Gerant Dashboard Views
@login_required(login_url='login')
def gerant_dashboard(request):
	"""Gerant dashboard showing clients and livreurs counts"""
	if request.user.role != 'gerant':
		return redirect('home')

	# Exclude soft-deleted users from dashboard counts
	clients_count = User.objects.filter(role='client', is_deleted=False).count()
	livreurs_count = User.objects.filter(role='livreur', is_deleted=False).count()
	
	# Get dish count from islem's menu (if running on port 8000)
	dishes_count = 0
	try:
		import requests
		response = requests.get('http://127.0.0.1:8000/api/dishes/count/', timeout=2)
		if response.status_code == 200:
			dishes_count = response.json().get('count', 0)
	except:
		# If islem server is not running or API doesn't exist, set to 0
		dishes_count = 0
	
	context = {
		'clients_count': clients_count,
		'livreurs_count': livreurs_count,
		'dishes_count': dishes_count,
	}
	return render(request, 'gerant/dashboard.html', context)


# Clients Management
@login_required(login_url='login')
def clients_list(request):
    """List all clients"""
    if request.user.role != 'gerant':
        return redirect('home')

    # Exclude soft-deleted clients from list
    clients = User.objects.filter(role='client', is_deleted=False).order_by('-date_joined')

    # Search/filter
    search_query = request.GET.get('search', '')
    if search_query:
        clients = clients.filter(
            Q(username__icontains=search_query) |
            Q(first_name__icontains=search_query) |
            Q(last_name__icontains=search_query) |
            Q(email__icontains=search_query)
        )

    context = {'clients': clients, 'search_query': search_query}
    return render(request, 'gerant/clients_list.html', context)


@login_required(login_url='login')
def client_detail(request, pk):
	"""Show client details"""
	if request.user.role != 'gerant':
		return redirect('home')
	
	client = get_object_or_404(User, pk=pk, role='client')
	return render(request, 'gerant/client_detail.html', {'client': client})


@login_required(login_url='login')
def client_create(request):
    if request.user.role != 'gerant':
        return redirect('home')

    if request.method == 'POST':
        form = ClientForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = 'client'
            user.set_password(form.cleaned_data['password1'])
            user.save()
            return redirect('clients_list')
    else:
        form = ClientForm()

    return render(request, 'gerant/client_form.html', {'form': form, 'action': 'Create'})

@login_required(login_url='login')
@login_required(login_url='login')
def client_update(request, pk):
    if request.user.role != 'gerant':
        return redirect('home')

    client = get_object_or_404(User, pk=pk, role='client')

    if request.method == 'POST':
        form = ClientForm(request.POST, instance=client)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = 'client'
            if form.cleaned_data.get('password1'):
                user.set_password(form.cleaned_data['password1'])
            user.save()
            return redirect('clients_list')
    else:
        form = ClientForm(instance=client)

    return render(request, 'gerant/client_form.html', {'form': form, 'action': 'Update', 'client': client})


@login_required(login_url='login')
def client_delete(request, pk):
    """Delete client"""
    if request.user.role != 'gerant':
        return redirect('home')

    client = get_object_or_404(User, pk=pk, role='client')

    if request.method == 'POST':
        # Soft delete: mark deleted and deactivate
        client.is_deleted = True
        client.is_active = False
        client.save()
        return redirect('clients_list')

    return render(request, 'gerant/client_confirm_delete.html', {'client': client})


# Livreurs Management
@login_required(login_url='login')
def livreurs_list(request):
    """List all livreurs"""
    if request.user.role != 'gerant':
        return redirect('home')

    # Exclude soft-deleted livreurs from list
    livreurs = User.objects.filter(role='livreur', is_deleted=False).order_by('-date_joined')

    # Search/filter
    search_query = request.GET.get('search', '')
    if search_query:
        livreurs = livreurs.filter(
            Q(username__icontains=search_query) |
            Q(first_name__icontains=search_query) |
            Q(last_name__icontains=search_query) |
            Q(email__icontains=search_query)
        )

    context = {'livreurs': livreurs, 'search_query': search_query}
    return render(request, 'gerant/livreurs_list.html', context)


@login_required(login_url='login')
def livreur_detail(request, pk):
	"""Show livreur details"""
	if request.user.role != 'gerant':
		return redirect('home')
	
	livreur = get_object_or_404(User, pk=pk, role='livreur')
	return render(request, 'gerant/livreur_detail.html', {'livreur': livreur})


@login_required(login_url='login')
@login_required(login_url='login')
def livreur_create(request):
    if request.user.role != 'gerant':
        return redirect('home')

    if request.method == 'POST':
        form = LivreurForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = 'livreur'
            user.set_password(form.cleaned_data['password1'])
            user.save()
            return redirect('livreurs_list')
    else:
        form = LivreurForm()

    return render(request, 'gerant/livreur_form.html', {'form': form, 'action': 'Create'})

@login_required(login_url='login')
@login_required(login_url='login')
def livreur_update(request, pk):
    if request.user.role != 'gerant':
        return redirect('home')

    livreur = get_object_or_404(User, pk=pk, role='livreur')

    if request.method == 'POST':
        form = LivreurForm(request.POST, instance=livreur)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = 'livreur'
            if form.cleaned_data.get('password1'):
                user.set_password(form.cleaned_data['password1'])
            user.save()
            return redirect('livreurs_list')
    else:
        form = LivreurForm(instance=livreur)

    return render(request, 'gerant/livreur_form.html', {'form': form, 'action': 'Update', 'livreur': livreur})

@login_required(login_url='login')
def livreur_delete(request, pk):
    """Delete livreur"""
    if request.user.role != 'gerant':
        return redirect('home')

    livreur = get_object_or_404(User, pk=pk, role='livreur')

    if request.method == 'POST':
        # Soft delete: mark deleted and deactivate
        livreur.is_deleted = True
        livreur.is_active = False
        livreur.save()
        return redirect('livreurs_list')

    return render(request, 'gerant/livreur_confirm_delete.html', {'livreur': livreur})
@login_required(login_url='login')
def client_delete_own_account(request):
    """Allow a client to delete ONLY their own account"""
    user = request.user

    # Only clients can delete themselves
    if user.role != 'client':
        return redirect('home')

    if request.method == 'POST':
        # Soft delete self and deactivate
        user.is_deleted = True
        user.is_active = False
        user.save()
        logout(request)
        return redirect('home')  # or a goodbye page

    return render(request, 'client_confirm_delete_own.html')


@login_required(login_url='login')
def client_toggle_active(request, pk):
    """Toggle client active/inactive status"""
    if request.user.role != 'gerant':
        return redirect('home')
    
    client = get_object_or_404(User, pk=pk, role='client')
    
    # Toggle the is_active status
    client.is_active = not client.is_active
    client.save()
    
    return redirect('client_detail', pk=pk)


@login_required(login_url='login')
def livreur_toggle_active(request, pk):
    """Toggle livreur active/inactive status"""
    if request.user.role != 'gerant':
        return redirect('home')
    
    livreur = get_object_or_404(User, pk=pk, role='livreur')
    
    # Toggle the is_active status
    livreur.is_active = not livreur.is_active
    livreur.save()
    
    return redirect('livreur_detail', pk=pk)


@login_required(login_url='login')
def export_clients_excel(request):
    """Export all clients to Excel file"""
    if request.user.role != 'gerant':
        return redirect('home')
    
    # Create workbook and worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"
    
    # Header styling
    header_fill = PatternFill(start_color="E67E22", end_color="E67E22", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Define headers
    headers = [
        "User ID", "Username", "First Name", "Last Name", "Email",
        "Phone Number", "Living Location", "Date of Birth", "Subscribed",
        "Active", "Deleted", "Date Joined"
    ]
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Get all clients
    # Include even soft-deleted clients
    clients = User.objects.filter(role='client').order_by('-date_joined')
    
    # Write data rows
    for row_num, client in enumerate(clients, 2):
        ws.cell(row=row_num, column=1).value = client.user_id
        ws.cell(row=row_num, column=2).value = client.username
        ws.cell(row=row_num, column=3).value = client.first_name
        ws.cell(row=row_num, column=4).value = client.last_name
        ws.cell(row=row_num, column=5).value = client.email
        ws.cell(row=row_num, column=6).value = client.phone_number or "N/A"
        ws.cell(row=row_num, column=7).value = client.living_location or "N/A"
        ws.cell(row=row_num, column=8).value = str(client.date_of_birth) if client.date_of_birth else "N/A"
        ws.cell(row=row_num, column=9).value = "Yes" if client.is_subscribed else "No"
        ws.cell(row=row_num, column=10).value = "Yes" if client.is_active else "No"
        ws.cell(row=row_num, column=11).value = "Yes" if getattr(client, 'is_deleted', False) else "No"
        ws.cell(row=row_num, column=12).value = client.date_joined.strftime("%Y-%m-%d %H:%M")
    
    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # Prepare response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"clients_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response


@login_required(login_url='login')
def export_livreurs_excel(request):
    """Export all livreurs to Excel file"""
    if request.user.role != 'gerant':
        return redirect('home')
    
    # Create workbook and worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Livreurs"
    
    # Header styling
    header_fill = PatternFill(start_color="E67E22", end_color="E67E22", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Define headers
    headers = [
        "User ID", "Username", "First Name", "Last Name", "Email",
        "Phone Number", "Vehicle Type", "Vehicle Number", "Vehicle Registration",
        "Insurance Number", "Bank Account", "Available", "Active", "Deleted", "Date Joined"
    ]
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # Get all livreurs
    # Include even soft-deleted livreurs
    livreurs = User.objects.filter(role='livreur').order_by('-date_joined')
    
    # Write data rows
    for row_num, livreur in enumerate(livreurs, 2):
        ws.cell(row=row_num, column=1).value = livreur.user_id
        ws.cell(row=row_num, column=2).value = livreur.username
        ws.cell(row=row_num, column=3).value = livreur.first_name
        ws.cell(row=row_num, column=4).value = livreur.last_name
        ws.cell(row=row_num, column=5).value = livreur.email
        ws.cell(row=row_num, column=6).value = livreur.phone_number or "N/A"
        ws.cell(row=row_num, column=7).value = livreur.vehicle_type or "N/A"
        ws.cell(row=row_num, column=8).value = livreur.vehicle_number or "N/A"
        ws.cell(row=row_num, column=9).value = livreur.vehicle_registration or "N/A"
        ws.cell(row=row_num, column=10).value = livreur.insurance_number or "N/A"
        ws.cell(row=row_num, column=11).value = livreur.bank_account or "N/A"
        ws.cell(row=row_num, column=12).value = "Yes" if livreur.is_available else "No"
        ws.cell(row=row_num, column=13).value = "Yes" if livreur.is_active else "No"
        ws.cell(row=row_num, column=14).value = "Yes" if getattr(livreur, 'is_deleted', False) else "No"
        ws.cell(row=row_num, column=15).value = livreur.date_joined.strftime("%Y-%m-%d %H:%M")
    
    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # Prepare response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"livreurs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response
    livreur.save()
    
    return redirect('livreur_detail', pk=pk)
